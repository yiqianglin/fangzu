"""
2026年5月1日收据 - 本地模式全流程脚本
- 上期读数：2026 年 3 月（3 月底抄表）
- 本期读数：2026 年 4 月（4 月底抄表）
- 收款日：2026 年 5 月 1 日
- 用量：4 月一个月用量

步骤：
1. 写 2026年3月/2026年3月水电表读数.xlsx
2. 写 2026年4月/2026年4月水电表读数.xlsx
3. 算 2026年5月/2026年5月应收租金.xlsx（含合计行）
4. 批量生成 16 份收据 xlsx → 2026年5月/收据/
5. 渲染 16 张收据图片 → 2026年5月/收据图片/

数据均硬编码，跑一次即可。
"""
import os
import shutil
from openpyxl import load_workbook, Workbook
from PIL import Image, ImageDraw, ImageFont

# ==================== 路径常量 ====================
ROOT = "/Users/linyiqiang/workspace/tencent/专利相关/每月租金生成/官田村"
COMMON = f"{ROOT}/通用"
TEMPLATE_HOUSE = f"{COMMON}/房屋说明.xlsx"
TEMPLATE_METER = f"{COMMON}/水电表读数-模板.xlsx"
TEMPLATE_RECEIPT = f"{COMMON}/收据模板.xlsx"

DIR_M3 = f"{ROOT}/2026年3月"
DIR_M4 = f"{ROOT}/2026年4月"
DIR_M5 = f"{ROOT}/2026年5月"
DIR_RECEIPT = f"{DIR_M5}/收据"
DIR_RECEIPT_IMG = f"{DIR_M5}/收据图片"

for d in (DIR_M3, DIR_M4, DIR_M5, DIR_RECEIPT, DIR_RECEIPT_IMG):
    os.makedirs(d, exist_ok=True)

WATER_PRICE = 6
ELEC_PRICE = 1

# ==================== 5 路 VLM 已确认的读数 ====================
# 17 间在租 + 401(仅水表，无人住，不入应收/收据，但月度表留痕)
# (房号, 水, 电, 备注)
READ_M3 = [
    ("101", 464, 1506, ""),
    ("102", 362, 5008, ""),
    ("103", 325, None, "电表合并入105"),
    ("105", 868, 5676, "含103合算"),
    ("106", 602, 7669, ""),
    ("201", 424, 465, ""),
    ("202", 484, 5261, ""),
    ("203", 560, 738, "刚租"),
    ("205", 453, 278, ""),
    ("206", 517, 9370, ""),
    ("207", 462, 4684, ""),
    ("301", 429, 7570, "刚租"),
    ("302", 480, 6603, ""),
    ("303", 504, 3603, ""),
    ("305", 704, 1133, ""),
    ("306", 333, 7470, "刚租"),
    ("307", 602, 5504, ""),
    ("401", 1641, None, "仅水表，无人租"),
]
READ_M4 = [
    ("101", 466, 1553, ""),
    ("102", 364, 5058, ""),
    ("103", 327, None, "电表合并入105"),
    ("105", 868, 5861, "含103合算（水表本月0用量，合并口径见应收）"),
    ("106", 602, 7673, "用量低，已与房东核对正常"),
    ("201", 425, 473, ""),
    ("202", 487, 5312, ""),
    ("203", 561, 764, ""),
    ("205", 456, 487, ""),
    ("206", 518, 9429, ""),
    ("207", 464, 4714, ""),
    ("301", 432, 7648, ""),
    ("302", 481, 6624, ""),
    ("303", 509, 3716, ""),
    ("305", 709, 1186, ""),
    ("306", 335, 7508, ""),
    ("307", 605, 5601, ""),
    ("401", 1641, None, "仅水表，无人租"),
]


def write_meter_xlsx(rows, out_path, sheet_name):
    """以 通用/水电表读数-模板.xlsx 为母版，写月度水电表"""
    src = load_workbook(TEMPLATE_METER)
    src.active.title = sheet_name
    ws = src.active
    # 模板第 1 行已是表头
    for r in rows:
        ws.append(r)
    src.save(out_path)
    print(f"✅ 写入月度水电表：{out_path}（{len(rows)} 行）")


# ==================== Step 1 & 2：写两份月度水电表 ====================
write_meter_xlsx(READ_M3, f"{DIR_M3}/2026年3月水电表读数.xlsx", "2026年3月水电表读数")
write_meter_xlsx(READ_M4, f"{DIR_M4}/2026年4月水电表读数.xlsx", "2026年4月水电表读数")


# ==================== Step 3：算 5 月应收租金 ====================
# 读房屋说明
hwb = load_workbook(TEMPLATE_HOUSE, data_only=True)
hws = hwb.active
HOUSE = {}  # 房号 -> dict
for row in hws.iter_rows(min_row=2, values_only=True):
    if not row or not row[0]:
        continue
    HOUSE[str(row[0])] = {
        "rented": bool(row[1]),
        "linked": row[2],
        "rent": row[3] or 0,
        "mgmt": row[4] or 0,
        "gas": row[5] or 0,
        "period_base": row[6] or "",
        "payee": row[7] or "张",
        "remark": row[8] or "",
    }

M3 = {r[0]: (r[1], r[2]) for r in READ_M3}  # 房号 -> (水, 电)
M4 = {r[0]: (r[1], r[2]) for r in READ_M4}


def map_period(base, target_month):
    """现算映射：基准文案 + 目标月偏移
    base: 如 '5月11日-6月11日' / '4月27日-5月27日' / '5月1日-5月30日'
    target_month: 目标本期收款日的月份（这里 = 5）
    """
    import re
    m = re.match(r'^\s*(\d+)月(\d+)日-(\d+)月(\d+)日\s*$', base)
    if not m:
        return base  # 异常文案，原样返回
    m1, d1, m2, d2 = (int(x) for x in m.groups())
    offset = target_month - m1
    nm1 = ((m1 - 1 + offset) % 12) + 1
    nm2 = ((m2 - 1 + offset) % 12) + 1
    return f"{nm1}月{d1}日-{nm2}月{d2}日"


def num_to_chinese(num):
    digits = ['零', '壹', '贰', '叁', '肆', '伍', '陆', '柒', '捌', '玖']
    units = ['', '拾', '佰', '仟', '万']
    if num == 0:
        return '零元整'
    result = ''
    s = str(int(num))
    L = len(s)
    for i, d in enumerate(s):
        digit = int(d)
        unit_idx = L - 1 - i
        if digit != 0:
            result += digits[digit] + units[unit_idx]
        else:
            if result and not result.endswith('零'):
                result += '零'
    result = result.rstrip('零')
    return result + '元整'


def split_amount_digits(amount):
    amt = int(amount)
    res = {'F': None, 'G': None, 'H': None, 'I': None, 'J': None}
    if amt >= 10000:
        res['F'] = (amt // 10000) % 10
    if amt >= 1000:
        res['G'] = (amt // 1000) % 10
    if amt >= 100:
        res['H'] = (amt // 100) % 10
    if amt >= 10:
        res['I'] = (amt // 10) % 10
    res['J'] = amt % 10
    return res


# 应收行
RECEIVABLE_HEADER = [
    "房号", "本期租期文案", "本期收款日",
    "房租", "管理费", "燃气费",
    "上月水表", "本月水表", "水表差", "水费",
    "上月电表", "本月电表", "电表差", "电费",
    "应收合计", "异常", "已收", "备注",
]

receivable_rows = []
TARGET_MONTH = 5
PAY_DATE = "2026年5月1日"

# 处理顺序（保持房屋说明里的顺序，但 103 跳过、401 跳过、合并到 105）
for room in ["101", "102", "105", "106",
             "201", "202", "203", "205", "206", "207",
             "301", "302", "303", "305", "306", "307"]:
    h = HOUSE[room]
    base = h["period_base"]
    period = map_period(base, TARGET_MONTH)
    rent = h["rent"]
    mgmt = h["mgmt"] or 0
    gas = h["gas"] or 0

    # 水表（105 特殊：103+105 合并）
    if room == "105":
        prev_w = M3["103"][0] + M3["105"][0]
        curr_w = M4["103"][0] + M4["105"][0]
        prev_e, curr_e = M3["105"][1], M4["105"][1]
        note = "含103合算（水表=103+105相加）"
    else:
        prev_w, prev_e = M3[room]
        curr_w, curr_e = M4[room]
        note = ""

    diff_w = curr_w - prev_w
    fee_w = diff_w * WATER_PRICE
    diff_e = curr_e - prev_e
    fee_e = diff_e * ELEC_PRICE

    abnormal = []
    if diff_w < 0:
        abnormal.append(f"水表倒退{diff_w}")
    if diff_e < 0:
        abnormal.append(f"电表倒退{diff_e}")
    abnormal_str = "; ".join(abnormal)

    total = rent + mgmt + gas + fee_w + fee_e

    receivable_rows.append([
        room, period, PAY_DATE,
        rent, mgmt, gas,
        prev_w, curr_w, diff_w, fee_w,
        prev_e, curr_e, diff_e, fee_e,
        total, abnormal_str, False, note,
    ])

# 合计行
sum_rent = sum(r[3] for r in receivable_rows)
sum_mgmt = sum(r[4] for r in receivable_rows)
sum_gas = sum(r[5] for r in receivable_rows)
sum_water_fee = sum(r[9] for r in receivable_rows)
sum_elec_fee = sum(r[13] for r in receivable_rows)
sum_total = sum(r[14] for r in receivable_rows)
n_abn = sum(1 for r in receivable_rows if r[15])

total_row = [
    "合计", "", "",
    sum_rent, sum_mgmt, sum_gas,
    "", "", "", sum_water_fee,
    "", "", "", sum_elec_fee,
    sum_total,
    f"含 {n_abn} 条异常" if n_abn else "",
    "",
    f"总水费 ¥{sum_water_fee} / 总电费 ¥{sum_elec_fee} / 总房租 ¥{sum_rent} / 总管理费 ¥{sum_mgmt} / 总燃气费 ¥{sum_gas} / 总应收 ¥{sum_total}",
]

# 写应收 xlsx
rwb = Workbook()
rws = rwb.active
rws.title = "2026年5月应收租金"
rws.append(RECEIVABLE_HEADER)
for r in receivable_rows:
    rws.append(r)
rws.append(total_row)
out_recv = f"{DIR_M5}/2026年5月应收租金.xlsx"
rwb.save(out_recv)
print(f"✅ 写入应收租金：{out_recv}（{len(receivable_rows)} 行 + 1 合计 / 总应收 ¥{sum_total}）")


# ==================== Step 4：批量生成收据 xlsx ====================
def gen_receipt(room, data):
    """data: receivable_rows 的一行"""
    period = data[1]  # 本期租期文案
    pay_date = data[2]
    rent = data[3]
    mgmt = data[4]
    gas = data[5]
    pw, cw, dw, fw = data[6], data[7], data[8], data[9]
    pe, ce, de, fe = data[10], data[11], data[12], data[13]
    total = data[14]

    out_xlsx = f"{DIR_RECEIPT}/{room}房_5月收据.xlsx"
    shutil.copy2(TEMPLATE_RECEIPT, out_xlsx)
    wb = load_workbook(out_xlsx)
    ws = wb.active
    ws.title = f"{room}房5月收据"

    ws['A2'] = f'用户名称：    {room}房'
    ws['H2'] = pay_date

    # 水费
    ws['B5'] = pw
    ws['C5'] = cw
    ws['D5'] = dw
    ws['E5'] = WATER_PRICE
    for col, v in split_amount_digits(fw).items():
        ws[f'{col}5'] = v

    # 电费
    ws['B6'] = pe
    ws['C6'] = ce
    ws['D6'] = de
    ws['E6'] = ELEC_PRICE
    for col, v in split_amount_digits(fe).items():
        ws[f'{col}6'] = v

    # 房租
    ws['B7'] = f"{period}房租"
    for col, v in split_amount_digits(rent).items():
        ws[f'{col}7'] = v

    # 卫生费/网费 留空金额
    for col in 'FGHIJ':
        ws[f'{col}8'] = None

    # 管理费 + 其它（燃气）
    ws['B9'] = mgmt if mgmt > 0 else None
    ws['D9'] = gas if gas > 0 else None
    row9 = mgmt + gas
    if row9 > 0:
        for col, v in split_amount_digits(row9).items():
            ws[f'{col}9'] = v
    else:
        for col in 'FGHIJ':
            ws[f'{col}9'] = None

    # 合计
    ws['A10'] = f'合计   人民币（大写）{num_to_chinese(total)}'
    ws['I10'] = f'(¥：{total}元)'

    wb.save(out_xlsx)
    return out_xlsx


print("\n📝 生成收据 xlsx：")
xlsx_paths = []
for r in receivable_rows:
    p = gen_receipt(r[0], r)
    xlsx_paths.append(p)
    print(f"  ✅ {r[0]}房 - 水{r[9]} 电{r[13]} 租{r[3]} 管{r[4]} 燃{r[5]} 合计 ¥{r[14]}")


# ==================== Step 5：渲染收据图片 ====================
FONT_PATH = "/System/Library/Fonts/STHeiti Medium.ttc"
IMG_WIDTH = 720
MARGIN_LEFT = 12
MARGIN_TOP = 8
BLACK = (0, 0, 0)
LIGHT_BG = (255, 255, 248)
LINE_COLOR = (60, 60, 60)
HEADER_BG = (245, 245, 235)


def load_fonts():
    return {
        "title": ImageFont.truetype(FONT_PATH, 20),
        "normal": ImageFont.truetype(FONT_PATH, 14),
        "small": ImageFont.truetype(FONT_PATH, 13),
        "bold": ImageFont.truetype(FONT_PATH, 15),
        "amount": ImageFont.truetype(FONT_PATH, 13),
    }


def draw_receipt(ws, output_path, fonts):
    MARGIN_BOTTOM = 8
    actual_height = MARGIN_TOP + 30 + 24 + (24 + 20 + 30 + 30 + 30 + 26 + 26 + 32 + 26) + MARGIN_BOTTOM
    img = Image.new("RGB", (IMG_WIDTH, actual_height), LIGHT_BG)
    draw = ImageDraw.Draw(img)
    x0 = MARGIN_LEFT
    y0 = MARGIN_TOP

    title = str(ws['A1'].value or "房租、水、电费（专用）收据")
    title_w = draw.textlength(title, font=fonts["title"])
    draw.text(((IMG_WIDTH - title_w) / 2, y0), title, fill=BLACK, font=fonts["title"])
    y0 += 30

    user_name = str(ws['A2'].value or "")
    date_str = str(ws['H2'].value or "")
    draw.text((x0, y0), user_name, fill=BLACK, font=fonts["normal"])
    date_w = draw.textlength(date_str, font=fonts["normal"])
    draw.text((IMG_WIDTH - MARGIN_LEFT - date_w, y0), date_str, fill=BLACK, font=fonts["normal"])
    y0 += 24

    table_width = IMG_WIDTH - 2 * MARGIN_LEFT
    col_widths = [80, 72, 72, 56, 56, 48, 48, 48, 48, 48, 0]
    col_widths[10] = table_width - sum(col_widths[:10])
    col_x = [x0]
    for w in col_widths:
        col_x.append(col_x[-1] + w)
    table_right = col_x[-1]

    row_heights = [24, 20, 30, 30, 30, 26, 26, 32, 26]

    def draw_row_bg(y, h, color):
        draw.rectangle([x0, y, table_right, y + h], fill=color)

    def draw_hline(y):
        draw.line([(x0, y), (table_right, y)], fill=LINE_COLOR, width=1)

    def draw_vlines(y, h):
        for i in range(len(col_x)):
            draw.line([(col_x[i], y), (col_x[i], y + h)], fill=LINE_COLOR, width=1)

    def text_center(text, col_idx, y, h, font_key="normal"):
        text = str(text) if text is not None else ""
        if not text:
            return
        f = fonts[font_key]
        tw = draw.textlength(text, font=f)
        tx = col_x[col_idx] + (col_widths[col_idx] - tw) / 2
        ty = y + (h - f.size) / 2 - 1
        draw.text((tx, ty), text, fill=BLACK, font=f)

    def text_center_span(text, col_start, col_end, y, h, font_key="normal"):
        text = str(text) if text is not None else ""
        if not text:
            return
        f = fonts[font_key]
        tw = draw.textlength(text, font=f)
        span_left = col_x[col_start]
        span_right = col_x[col_end] + col_widths[col_end]
        tx = span_left + (span_right - span_left - tw) / 2
        ty = y + (h - f.size) / 2 - 1
        draw.text((tx, ty), text, fill=BLACK, font=f)

    def text_left(text, x, y, h, font_key="normal"):
        f = fonts[font_key]
        ty = y + (h - f.size) / 2 - 1
        draw.text((x + 4, ty), str(text), fill=BLACK, font=f)

    def text_right(text, x_right, y, h, font_key="normal"):
        f = fonts[font_key]
        tw = draw.textlength(str(text), font=f)
        ty = y + (h - f.size) / 2 - 1
        draw.text((x_right - tw - 4, ty), str(text), fill=BLACK, font=f)

    # 表头
    y = y0
    h = row_heights[0]
    h2 = row_heights[1]
    total_header_h = h + h2
    draw_row_bg(y, total_header_h, HEADER_BG)
    draw_hline(y)

    text_center("项 目", 0, y, total_header_h, "bold")
    text_center("上 月", 1, y, total_header_h, "bold")
    text_center("本 月", 2, y, total_header_h, "bold")
    text_center("实 用", 3, y, total_header_h, "bold")
    text_center("单 价", 4, y, total_header_h, "bold")
    text_center_span("金    额", 5, 9, y, h, "bold")
    text_center("备注", 10, y, total_header_h, "bold")

    for i in [0, 1, 2, 3, 4, 5, 10, 11]:
        draw.line([(col_x[i], y), (col_x[i], y + total_header_h)], fill=LINE_COLOR, width=1)
    draw.line([(col_x[5], y + h), (col_x[10], y + h)], fill=LINE_COLOR, width=1)

    y += h
    h = h2
    for i, label in enumerate(['万', '千', '百', '十', '元']):
        text_center(label, 5 + i, y, h, "small")
    for i in range(6, 10):
        draw.line([(col_x[i], y), (col_x[i], y + h)], fill=LINE_COLOR, width=1)
    draw_hline(y + h)

    def draw_data_row(y, h, label, ws_row):
        draw_hline(y)
        text_center(label, 0, y, h, "normal")
        for ci, col_letter in enumerate(['B', 'C', 'D', 'E'], 1):
            val = ws[f'{col_letter}{ws_row}'].value
            if val is not None:
                text_center(str(val), ci, y, h, "normal")
        for ci, col_letter in enumerate(['F', 'G', 'H', 'I', 'J'], 5):
            val = ws[f'{col_letter}{ws_row}'].value
            if val is not None:
                text_center(str(val), ci, y, h, "amount")
        val = ws[f'K{ws_row}'].value
        if val is not None:
            text_center(str(val), 10, y, h, "small")
        draw_vlines(y, h)

    y += h
    h = row_heights[2]
    draw_data_row(y, h, "水费(吨)", 5)

    y += h
    h = row_heights[3]
    draw_data_row(y, h, "电费(度)", 6)

    # 房租
    y += h
    h = row_heights[4]
    draw_hline(y)
    text_center("房 租", 0, y, h, "normal")
    rent_period = ws['B7'].value
    if rent_period:
        text_center_span(str(rent_period), 1, 4, y, h, "normal")
    for ci, col_letter in enumerate(['F', 'G', 'H', 'I', 'J'], 5):
        val = ws[f'{col_letter}7'].value
        if val is not None:
            text_center(str(val), ci, y, h, "amount")
    merge_skip = {2, 3, 4}
    for i in range(len(col_x)):
        if i not in merge_skip:
            draw.line([(col_x[i], y), (col_x[i], y + h)], fill=LINE_COLOR, width=1)

    # 卫生费/网费
    y += h
    h = row_heights[5]
    draw_hline(y)
    text_center("卫生费", 0, y, h, "small")
    text_center("网 费", 2, y, h, "small")
    draw_vlines(y, h)

    # 管理费 + 其它
    y += h
    h = row_heights[6]
    draw_hline(y)
    text_center("管理费", 0, y, h, "small")
    val_b9 = ws['B9'].value
    if val_b9:
        text_center(str(val_b9), 1, y, h, "normal")
    text_center("其 它", 2, y, h, "small")
    val_d9 = ws['D9'].value
    if val_d9:
        text_center(str(val_d9), 3, y, h, "normal")
    for ci, col_letter in enumerate(['F', 'G', 'H', 'I', 'J'], 5):
        val = ws[f'{col_letter}9'].value
        if val is not None:
            text_center(str(val), ci, y, h, "amount")
    draw_vlines(y, h)

    # 合计
    y += h
    h = row_heights[7]
    draw_hline(y)
    total_text = str(ws['A10'].value or "")
    text_left(total_text, x0, y, h, "normal")
    amount_text = str(ws['I10'].value or "")
    text_right(amount_text, table_right - 4, y, h, "bold")
    draw.line([(x0, y), (x0, y + h)], fill=LINE_COLOR, width=1)
    draw.line([(table_right, y), (table_right, y + h)], fill=LINE_COLOR, width=1)
    draw_hline(y + h)

    # 开票人
    y += h
    h = row_heights[8]
    issuer = str(ws['A11'].value or "开票人：")
    text_left(issuer, x0, y, h, "normal")
    receiver = str(ws['H11'].value or "收款人：张")
    text_right(receiver, table_right, y, h, "normal")

    img.save(output_path, "JPEG", quality=95)


print("\n🖼️  渲染收据图片：")
fonts = load_fonts()
for xlsx_path in xlsx_paths:
    fname = os.path.basename(xlsx_path)
    img_name = fname.replace('.xlsx', '.jpg')
    img_path = f"{DIR_RECEIPT_IMG}/{img_name}"
    wb = load_workbook(xlsx_path)
    ws = wb.active
    draw_receipt(ws, img_path, fonts)
    wb.close()
    print(f"  ✅ {img_name}")


# ==================== 汇总 ====================
print(f"\n🎉 全部完成！")
print(f"  - 月度水电表：{DIR_M3}/2026年3月水电表读数.xlsx")
print(f"  - 月度水电表：{DIR_M4}/2026年4月水电表读数.xlsx")
print(f"  - 应收租金：{out_recv}")
print(f"  - 收据 xlsx：{DIR_RECEIPT}（{len(xlsx_paths)} 份）")
print(f"  - 收据图片：{DIR_RECEIPT_IMG}（{len(xlsx_paths)} 张）")
print(f"  - 总应收：¥{sum_total}（房租 ¥{sum_rent} / 管理 ¥{sum_mgmt} / 燃气 ¥{sum_gas} / 水费 ¥{sum_water_fee} / 电费 ¥{sum_elec_fee}）")
