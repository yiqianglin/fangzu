"""
清理「房屋说明-{村}.xlsx」到 v2 干净状态：

1. 主 sheet 只保留档案字段：房号 / 是否在租 / 打通房 / 收款人 / 备注 / 水单价 / 水保底度数
   删除：房租 / 管理费 / 燃气费 / 当前租期文案 / 租客姓名 / 下月房租
   （这些数据全部由「租金月历」sheet 接管）

2. 「租金月历」sheet：
   - 「初始基准」行的 生效起月 从 2024-01 改成 2026-01
   - 「自然月」改写成 「1日-30日」

3. 保留单元格样式（边框 / 字体 / 对齐 / 字号），openpyxl 的 cell.font cell.border 等会自动跟随。
   方案：直接用 openpyxl 操作，删列用 ws.delete_cols，改值用 ws.cell.value = ...
"""
from __future__ import annotations
import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent  # 通用/
KEEP_MAIN_COLS = ["房号", "是否在租", "打通房", "收款人", "备注", "水单价", "水保底度数"]


def clean_main_sheet(ws):
    """主 sheet：只保留 KEEP_MAIN_COLS。其余列全部删除。"""
    hdr = [c.value for c in ws[1]]
    # 找出要删除的列（按从右到左删，避免下标位移）
    cols_to_delete = []
    for i, name in enumerate(hdr, 1):
        if name not in KEEP_MAIN_COLS:
            cols_to_delete.append(i)
    cols_to_delete.sort(reverse=True)
    for ci in cols_to_delete:
        ws.delete_cols(ci, 1)
    print(f"  主 sheet：删除 {len(cols_to_delete)} 列，保留 {[c.value for c in ws[1]]}")


def clean_calendar_sheet(ws):
    """租金月历：
    - 「初始基准」行 (备注列以"初始基准"开头) 的 生效起月 改 2026-01
    - 「自然月」→「1日-30日」
    """
    hdr = [c.value for c in ws[1]]
    idx = {n: i for i, n in enumerate(hdr) if n}
    col_ym = idx["生效起月"]
    col_period = idx["租约日段"]
    col_remark = idx["备注"]

    n_ym, n_period = 0, 0
    for row in ws.iter_rows(min_row=2):
        ym_cell = row[col_ym]
        period_cell = row[col_period]
        remark_cell = row[col_remark]

        # 改 生效起月
        if ym_cell.value == "2024-01" and remark_cell.value and str(remark_cell.value).startswith("初始基准"):
            ym_cell.value = "2026-01"
            n_ym += 1

        # 改 租约日段
        if period_cell.value == "自然月":
            period_cell.value = "1日-30日"
            n_period += 1
    print(f"  租金月历：改 生效起月 {n_ym} 行 / 改「自然月→1日-30日」{n_period} 行")


def main():
    for v in ["官田村", "变电站"]:
        p = ROOT / f"房屋说明-{v}.xlsx"
        if not p.exists():
            print(f"⏭  跳过：{p} 不存在")
            continue
        print(f"\n=== {v} ({p.name}) ===")
        wb = load_workbook(p)
        # 主 sheet
        ws_main = wb.active
        clean_main_sheet(ws_main)
        # 月历
        if "租金月历" in wb.sheetnames:
            clean_calendar_sheet(wb["租金月历"])
        else:
            print(f"  ⚠️ {v} 缺少「租金月历」sheet，跳过")
        wb.save(p)
        wb.close()
        print(f"  ✅ 已写回 {p}")


if __name__ == "__main__":
    main()
