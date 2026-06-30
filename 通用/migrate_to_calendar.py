# -*- coding: utf-8 -*-
"""
一次性迁移脚本：把「房屋说明.xlsx」改造成支持月历的结构。

执行后会：
1. 在原 xlsx 里新增一个 sheet：`租金月历`
2. 列：房号 / 生效起月 / 租客姓名 / 房租 / 管理费 / 燃气费 / 租约日段 / 备注
3. 第一行 = 初始基准（生效起月 2024-01）
4. 已知涨租的房号 自动追加第二行（生效起月 2026-07）

旧字段（房租/管理费/燃气费/当前租期文案/租客姓名/下月房租）保留不动，
让用户回查时仍能看到原始信息；脚本以后只读 `租金月历` + `房屋说明`(精简字段)。

幂等：如果已存在 `租金月历` sheet，先删除再重建。
"""
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

BASE_DIR = "/Users/linyiqiang/workspace/tencent/专利相关/每月租金生成/通用"

# 已知 2026-07 起涨租（用户确认：当前 6 月按旧价，7 月起新价）
RENT_CHANGES_2026_07 = {
    "官田村": {
        "206": 720,
        "305": 680,
        "307": 730,
    },
    "变电站": {},  # 变电站当前没有「下月房租」字段，无涨租
}


def parse_lease_pattern(text):
    """
    解析旧的「当前租期文案」，输出抽象「租约日段」。

    规则：
    - "5月1日-5月30日" / "5月01日-5月30日" → "自然月"
    - "5月10日-6月10日" → "10日-次月10日"
    - "5月X日-6月X日"（X≠1）→ "X日-次月X日"
    - 含 "与XXX合算" → 沿用原文
    - None / 空 → None
    """
    if not text:
        return None
    s = str(text).strip()
    if "合算" in s or "见" in s:
        return s  # 合并房号原样保留
    # 匹配 "5月X日-Y月Z日"
    m = re.match(r"^(\d+)月(\d+)日\s*[-–~]\s*(\d+)月(\d+)日$", s)
    if not m:
        return s  # 不认识的格式原样保留
    m1, d1, m2, d2 = int(m.group(1)), int(m.group(2)), int(m.group(3)), int(m.group(4))
    if m1 == m2:
        # 同月，且 d1=1, d2 in (28,29,30,31) → 自然月
        if d1 == 1 and d2 >= 28:
            return "自然月"
        return f"{d1}日-{d2}日"  # 罕见情况
    else:
        # 跨月：m2 = m1 + 1，且 d1 == d2 → 标准跨月制
        if d1 == d2:
            return f"{d1}日-次月{d1}日"
        # 跨月但起止日不同（罕见），原样保留信息
        return f"{d1}日-次月{d2}日"


def migrate_one(name):
    path = f"{BASE_DIR}/房屋说明-{name}.xlsx"
    print(f"\n=== 迁移 {path} ===")
    wb = openpyxl.load_workbook(path)
    src = wb["房屋说明"]

    # 读取表头位置
    header = [c.value for c in src[1]]
    col = {h: i for i, h in enumerate(header)}
    print(f"  原表头: {header}")

    def get(row, key, default=None):
        i = col.get(key)
        if i is None:
            return default
        v = row[i]
        return v if v is not None else default

    # 删除已存在的「租金月历」（幂等）
    if "租金月历" in wb.sheetnames:
        del wb["租金月历"]
        print("  已删除旧的「租金月历」sheet")

    cal = wb.create_sheet("租金月历")
    cal_header = ["房号", "生效起月", "租客姓名", "房租", "管理费", "燃气费", "租约日段", "备注"]
    cal.append(cal_header)
    # 表头加粗 + 填充色
    for c in cal[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFE599")
        c.alignment = Alignment(horizontal="center")

    rows_added = 0
    rent_changes = RENT_CHANGES_2026_07.get(name, {})

    # 遍历 房屋说明 数据行
    for row in src.iter_rows(min_row=2, values_only=True):
        rno_raw = row[col["房号"]]
        if rno_raw is None:
            continue
        rno = str(rno_raw)
        in_rent = get(row, "是否在租")
        rent = get(row, "房租")
        mgmt = get(row, "管理费")
        gas = get(row, "燃气费")
        lease_text = get(row, "当前租期文案")
        tenant = get(row, "租客姓名")
        remark = get(row, "备注")
        next_rent = get(row, "下月房租")  # 仅官田村有

        # 不在租 / 没房租，跳过（401/402/501/601 这种）
        if in_rent is False or (rent is None and not lease_text):
            print(f"  跳过 {rno}（不在租或无数据）")
            continue

        lease_pattern = parse_lease_pattern(lease_text)

        # 第一行：初始基准
        cal.append([
            rno,
            "2024-01",
            tenant,
            rent,
            mgmt,
            gas,
            lease_pattern,
            "初始基准" + (f"（{remark}）" if remark else ""),
        ])
        rows_added += 1

        # 第二行（如有涨租）
        if rno in rent_changes:
            new_rent = rent_changes[rno]
            cal.append([
                rno,
                "2026-07",
                tenant,
                new_rent,
                mgmt,
                gas,
                lease_pattern,
                f"涨租：{rent} → {new_rent}",
            ])
            rows_added += 1
            print(f"  ✓ {rno} 追加涨租行：{rent} → {new_rent}（生效 2026-07）")

    # 列宽
    widths = [8, 12, 12, 8, 8, 8, 18, 28]
    for i, w in enumerate(widths, 1):
        cal.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    wb.save(path)
    print(f"  ✓ 已写入 {rows_added} 行到「租金月历」sheet")


if __name__ == "__main__":
    for name in ["官田村", "变电站"]:
        migrate_one(name)
    print("\n全部迁移完成。")
