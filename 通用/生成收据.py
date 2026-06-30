# -*- coding: utf-8 -*-
"""
【固化版】批量生成收据 xlsx —— 数据从「YYYY年M月应收租金.xlsx」读，无 hardcode。

用法：
    python 每月租金生成/通用/生成收据.py --year 2026 --month 6
    python 每月租金生成/通用/生成收据.py --year 2026 --month 6 --village 变电站   # 指定小区（默认官田村）

输入：
    {ROOT}/{village}/{year}年{month}月/{year}年{month}月应收租金.xlsx   # 第 2 步产出
    {ROOT}/通用/房屋说明-{village}.xlsx                                 # 静态档案（取水单价 + 租金月历 sheet）
    {ROOT}/通用/收据模板.xlsx                                           # 收据模板

输出：
    {ROOT}/{village}/{year}年{month}月/收据/{房号}房_{month}月收据.xlsx

铁律（写进 SKILL.md）：
- A2 用户名称 = `{房号}房  {姓名}`（房号 + 两个空格 + 姓名），姓名缺失要警告
- K7 下月预告：对比 (当月房租) vs (下月房租)，不一致才写 `下月（X月）房租改为XXX元`
  - 数据源：「房屋说明.租金月历」sheet → lookup_calendar(rno, year, month)
- K5 / K6 / K9 / K10 一律不写下月预告
- 「合计」行跳过，不出收据
- 103 不出收据（在第 2 步已合并到 105）
"""
import argparse
import os
import shutil
import sys
from openpyxl import load_workbook


# ==================== 工具函数 ====================
def num_to_chinese(num):
    digits = ['零', '壹', '贰', '叁', '肆', '伍', '陆', '柒', '捌', '玖']
    units = ['', '拾', '佰', '仟', '万']
    if num == 0:
        return '零元整'
    result = ''
    s = str(int(num))
    L = len(s)
    for i, d in enumerate(s):
        digit = int(d)
        unit_idx = L - 1 - i
        if digit != 0:
            result += digits[digit] + units[unit_idx]
        else:
            if result and not result.endswith('零'):
                result += '零'
    result = result.rstrip('零')
    return result + '元整'


def split_amount_digits(amount):
    """金额拆分到 万千百十元（F~J 五位）"""
    amt = int(amount)
    res = {'F': None, 'G': None, 'H': None, 'I': None, 'J': None}
    if amt >= 10000:
        res['F'] = (amt // 10000) % 10
    if amt >= 1000:
        res['G'] = (amt // 1000) % 10
    if amt >= 100:
        res['H'] = (amt // 100) % 10
    if amt >= 10:
        res['I'] = (amt // 10) % 10
    res['J'] = amt % 10
    return res


def load_house(template_house):
    """读「房屋说明-{村}.xlsx」主 sheet 拿水单价（姓名/房租改用月历）"""
    wb = load_workbook(template_house, data_only=True)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    idx = {name: hdr.index(name) for name in hdr if name}
    house = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        rno = str(row[0]).strip()
        house[rno] = {
            "water_price": row[idx["水单价"]] if idx.get("水单价") is not None else None,
        }
    wb.close()
    return house


def load_calendar(template_house):
    """读「房屋说明-{村}.xlsx → 租金月历」sheet。
    返回 {rno: [{'year':..., 'month':..., 'name':..., 'rent':..., 'mgmt':..., 'gas':..., 'period':...}, ...]}
    每个房号的列表已按 (year,month) 升序排序。
    """
    wb = load_workbook(template_house, data_only=True)
    if "租金月历" not in wb.sheetnames:
        wb.close()
        raise RuntimeError(f"❌ {template_house} 缺少「租金月历」sheet，请先跑迁移脚本 migrate_to_calendar.py")
    ws = wb["租金月历"]
    hdr = [c.value for c in ws[1]]
    idx = {name: hdr.index(name) for name in hdr if name}
    cal = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        rno = str(row[0]).strip()
        ym_raw = row[idx["生效起月"]] if idx.get("生效起月") is not None else None
        if not ym_raw:
            continue
        ym_s = str(ym_raw).strip()
        try:
            y_str, m_str = ym_s.split("-")
            y, m = int(y_str), int(m_str)
        except (ValueError, AttributeError):
            print(f"  ⚠️  租金月历格式异常：{rno} 生效起月={ym_raw!r}（应为 YYYY-MM）", file=sys.stderr)
            continue
        cal.setdefault(rno, []).append({
            "year": y,
            "month": m,
            "name": (row[idx["租客姓名"]] or "").strip() if idx.get("租客姓名") is not None and row[idx["租客姓名"]] else "",
            "rent": row[idx["房租"]] if idx.get("房租") is not None else None,
            "mgmt": row[idx["管理费"]] if idx.get("管理费") is not None else None,
            "gas": row[idx["燃气费"]] if idx.get("燃气费") is not None else None,
            "period": (row[idx["租约日段"]] or "").strip() if idx.get("租约日段") is not None and row[idx["租约日段"]] else "",
        })
    wb.close()
    for rno in cal:
        cal[rno].sort(key=lambda x: (x["year"], x["month"]))
    return cal


def lookup_calendar(cal, rno, year, month):
    """查 (rno, year, month) 当时生效的月历行：取 (生效年,生效月) <= (year,month) 中最近的一行。
    找不到返回 None。"""
    rows = cal.get(str(rno).strip(), [])
    target = (year, month)
    hit = None
    for r in rows:
        if (r["year"], r["month"]) <= target:
            hit = r
        else:
            break
    return hit


def load_receivable(receivable_xlsx):
    """读「YYYY年M月应收租金.xlsx」"""
    wb = load_workbook(receivable_xlsx, data_only=True)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    idx = {name: hdr.index(name) for name in hdr if name}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        rno = str(row[0]).strip()
        # 只接受纯数字房号；「合计」/汇总文案/空行一律跳过
        if not rno.isdigit():
            continue
        rows.append(row)
    wb.close()
    return idx, rows


# ==================== 单张收据生成 ====================
def gen_receipt(rno_data, ridx, house, cal, target_year, target_month,
                water_price_default, elec_price, template_receipt, dir_receipt):
    """生成 1 张收据 xlsx，返回 (路径, 房号, 各金额)"""
    rno = str(rno_data[ridx["房号"]]).strip()
    period = rno_data[ridx["本期租期文案"]]
    pay_date = rno_data[ridx["本期收款日"]]
    rent = rno_data[ridx["房租"]] or 0
    mgmt = rno_data[ridx["管理费"]] or 0
    gas = rno_data[ridx["燃气费"]] or 0
    pw = rno_data[ridx["上月水表"]]
    cw = rno_data[ridx["本月水表"]]
    dw = rno_data[ridx["水表差"]]
    fw = rno_data[ridx["水费"]] or 0
    pe = rno_data[ridx["上月电表"]]
    ce = rno_data[ridx["本月电表"]]
    de = rno_data[ridx["电表差"]]
    fe = rno_data[ridx["电费"]] or 0
    total = rno_data[ridx["应收合计"]] or 0

    h = house.get(rno, {})
    water_price = h.get("water_price") or water_price_default

    # 月历查 当月 + 下月
    cur = lookup_calendar(cal, rno, target_year, target_month)
    next_year = target_year + 1 if target_month == 12 else target_year
    next_month = 1 if target_month == 12 else target_month + 1
    nxt = lookup_calendar(cal, rno, next_year, next_month)

    name = (cur or {}).get("name", "")

    out_xlsx = os.path.join(dir_receipt, f"{rno}房_{target_month}月收据.xlsx")
    shutil.copy2(template_receipt, out_xlsx)
    wb = load_workbook(out_xlsx)
    ws = wb.active
    ws.title = f"{rno}房{target_month}月收据"

    # A2 用户名称
    if name:
        ws['A2'] = f'用户名称：    {rno}房  {name}'
    else:
        ws['A2'] = f'用户名称：    {rno}房'
        print(f"  ⚠️  {rno}房：租金月历未填租客姓名，请到「房屋说明.租金月历」补全", file=sys.stderr)
    ws['H2'] = pay_date

    # 水费
    ws['B5'] = pw
    ws['C5'] = cw
    ws['D5'] = dw
    ws['E5'] = water_price
    for col, v in split_amount_digits(fw).items():
        ws[f'{col}5'] = v

    # 电费
    ws['B6'] = pe
    ws['C6'] = ce
    ws['D6'] = de
    ws['E6'] = elec_price
    for col, v in split_amount_digits(fe).items():
        ws[f'{col}6'] = v

    # 房租
    ws['B7'] = f"{period}房租"
    for col, v in split_amount_digits(rent).items():
        ws[f'{col}7'] = v

    # K7 下月预告（对比 当月 vs 下月房租，不一致才写）
    cur_rent = (cur or {}).get("rent")
    nxt_rent = (nxt or {}).get("rent")
    if cur_rent is not None and nxt_rent is not None and int(cur_rent) != int(nxt_rent):
        ws['K7'] = f"下月（{next_month}月）房租改为{int(nxt_rent)}元"
    else:
        ws['K7'] = None

    # 卫生费/网费 留空
    for col in 'FGHIJ':
        ws[f'{col}8'] = None

    # 管理费 + 其它（燃气）
    ws['B9'] = mgmt if mgmt > 0 else None
    ws['D9'] = gas if gas > 0 else None
    row9 = (mgmt or 0) + (gas or 0)
    if row9 > 0:
        for col, v in split_amount_digits(row9).items():
            ws[f'{col}9'] = v
    else:
        for col in 'FGHIJ':
            ws[f'{col}9'] = None

    # 合计
    ws['A10'] = f'合计   人民币（大写）{num_to_chinese(total)}'
    ws['I10'] = f'(¥：{total}元)'

    wb.save(out_xlsx)
    return out_xlsx, rno, fw, fe, rent, mgmt, gas, total


# ==================== 主流程 ====================
def main():
    ap = argparse.ArgumentParser(description="批量生成收据 xlsx（数据从应收租金 xlsx 读）")
    ap.add_argument("--year", type=int, required=True, help="年份，如 2026")
    ap.add_argument("--month", type=int, required=True, help="月份（1-12），如 6")
    ap.add_argument("--village", default="官田村", help="小区名，默认 官田村")
    ap.add_argument("--root", default=None, help="根目录，默认 = 脚本所在的「通用」的父目录")
    ap.add_argument("--water-price", type=float, default=6, help="水单价默认值（元/吨），默认 6")
    ap.add_argument("--elec-price", type=float, default=1, help="电单价（元/度），默认 1")
    args = ap.parse_args()

    # 路径解析：脚本在 {root}/通用/ 下，root 默认推断为脚本所在目录的父目录
    script_dir = os.path.dirname(os.path.abspath(__file__))
    if args.root:
        root = args.root
    else:
        # 通用/ → ROOT
        root = os.path.dirname(script_dir)

    common = os.path.join(root, "通用")
    village_dir = os.path.join(root, args.village)
    template_house = os.path.join(common, f"房屋说明-{args.village}.xlsx")
    template_receipt = os.path.join(common, "收据模板.xlsx")

    month_dir = os.path.join(village_dir, f"{args.year}年{args.month}月")
    receivable_xlsx = os.path.join(month_dir, f"{args.year}年{args.month}月应收租金.xlsx")
    dir_receipt = os.path.join(month_dir, "收据")

    # 前置校验
    for p, desc in [
        (template_house, "房屋说明"),
        (template_receipt, "收据模板"),
        (receivable_xlsx, "应收租金 xlsx（请先跑第 2 步）"),
    ]:
        if not os.path.exists(p):
            print(f"❌ 找不到{desc}：{p}", file=sys.stderr)
            sys.exit(1)

    os.makedirs(dir_receipt, exist_ok=True)

    print(f"📂 小区：{args.village}  目标月：{args.year} 年 {args.month} 月")
    print(f"📥 应收租金：{receivable_xlsx}")
    print(f"📤 收据输出：{dir_receipt}")
    print()

    house = load_house(template_house)
    cal = load_calendar(template_house)
    ridx, rows = load_receivable(receivable_xlsx)
    print(f"📊 应收租金 {len(rows)} 行（已剔除合计）")
    print(f"📅 租金月历 {sum(len(v) for v in cal.values())} 行（覆盖 {len(cal)} 个房号）")

    print("\n📝 生成收据 xlsx：")
    results = []
    for r in rows:
        info = gen_receipt(r, ridx, house, cal, args.year, args.month,
                           args.water_price, args.elec_price,
                           template_receipt, dir_receipt)
        results.append(info)
        _, rno, fw, fe, rent, mgmt, gas, total = info
        print(f"  ✅ {rno}房 - 水{fw} 电{fe} 租{rent} 管{mgmt} 燃{gas} 合计 ¥{total}")

    # 汇总
    sum_total = sum(x[7] for x in results)
    sum_rent = sum(x[4] for x in results)
    sum_mgmt = sum(x[5] for x in results)
    sum_gas = sum(x[6] for x in results)
    sum_w = sum(x[2] for x in results)
    sum_e = sum(x[3] for x in results)
    print(f"\n🎉 收据 xlsx 共 {len(results)} 份")
    print(f"  总应收 ¥{sum_total}（房租 ¥{sum_rent} / 管理 ¥{sum_mgmt} / 燃气 ¥{sum_gas} / 水费 ¥{sum_w} / 电费 ¥{sum_e}）")
    print(f"  → 下一步：python 每月租金生成/通用/导出收据图片.py --year {args.year} --month {args.month} --village {args.village}")


if __name__ == "__main__":
    main()
